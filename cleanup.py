#cleanup.py
import pandas as pd
import io
import re

def read_docx_tables(file_bytes: bytes):
    """Extracts tables from a Word document (.docx)."""
    try:
        from docx import Document
    except Exception as e:
        raise ImportError("python-docx is required to read .docx files. Install with `pip install python-docx`") from e

    doc = Document(io.BytesIO(file_bytes))
    tables = []
    # --- Extract actual tables ---
    for table in doc.tables:
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        if data:
            tables.append(pd.DataFrame(data))

    
    paragraph_data = extract_paragraph_data(doc)
    if paragraph_data is not None:
        tables.append(paragraph_data)

    return tables






def extract_paragraph_data(doc):
    """Detect paragraph data like Name:Abi, Age:20, City:Madurai."""
    # Split paragraphs and keep even empty lines to detect record breaks
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    lines = [p for p in paragraphs if p or p == ""]

    records = []
    current = {}

    # Detect available keys dynamically (preserve order of first appearance)
    key_pattern = re.compile(r"^([\w\s]+?)\s*:")
    all_keys_ordered = []

    # First pass – collect keys in first-seen order
    for line in lines:
        m = key_pattern.match(line)
        if m:
            k = m.group(1).strip()
            if k not in all_keys_ordered:
                all_keys_ordered.append(k)

    # Second pass – build records
    for line in lines:
        if not line.strip():  # blank line = end of one record
            if current:
                records.append(current)
                current = {}
            continue

        if ":" in line:
            k, v = line.split(":", 1)
            k, v = k.strip(), v.strip()
            current[k] = v

    if current:
        records.append(current)

    # If no structured key-value data found
    if not records or len(all_keys_ordered) < 1:
        return None

    # Keep only complete records where all keys are present and non-empty
    complete_records = []
    for r in records:
        if all((k in r and str(r[k]).strip() != "") for k in all_keys_ordered):
            complete_records.append(r)

    if not complete_records:
        return None

    # Convert to DataFrame using ordered columns
    df = pd.DataFrame(complete_records, columns=all_keys_ordered)
    # Title-case column names
    df.columns = [c.strip().title() for c in df.columns]
    return df